"""Shared helper for writing validation-sample output as .xlsx.

Used by every country script so all deliverables share the same
formatting behaviour (frozen header, sensible column widths, and
forced text formatting on ID-like columns so leading zeros / long
digit strings never get silently reinterpreted as numbers).
"""
import re

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

# XML 1.0 disallows these control characters outright - openpyxl raises
# IllegalCharacterError rather than writing them (confirmed on a real scraped
# page title containing a literal \x13: "Geddington CofE Primary School
# \x13 Excellence..."). These are junk bytes from the source page's own
# markup/encoding, not meaningful content, so stripping is correct rather
# than a workaround.
ILLEGAL_XLSX_CHARS_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _clean(value):
    if isinstance(value, str):
        return ILLEGAL_XLSX_CHARS_RE.sub("", value)
    return value


def write_xlsx(rows, fieldnames, out_path, sheet_name="Sheet1", text_columns=None):
    """
    rows: list of dicts keyed by fieldnames
    fieldnames: ordered list of column names
    text_columns: set of column names to force to text format (e.g. IDs, phone numbers, postcodes)
    """
    text_columns = text_columns or set()
    text_col_idx = {i for i, h in enumerate(fieldnames) if h in text_columns}

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    ws.append(fieldnames)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        ws.append([_clean(row.get(h, "")) for h in fieldnames])

    last_row = len(rows) + 1
    for idx in text_col_idx:
        col_letter = get_column_letter(idx + 1)
        for r in range(2, last_row + 1):
            ws[f"{col_letter}{r}"].number_format = "@"

    for i, h in enumerate(fieldnames):
        values = [str(row.get(h, "")) for row in rows]
        max_len = max([len(h)] + [len(v) for v in values]) if values else len(h)
        ws.column_dimensions[get_column_letter(i + 1)].width = min(max(max_len + 2, 10), 45)

    ws.freeze_panes = "A2"
    wb.save(out_path)
    return out_path
